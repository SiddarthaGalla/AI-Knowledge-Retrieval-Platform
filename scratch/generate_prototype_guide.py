import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_prototype_guide_files():
    # -------------------------------------------------------------------------
    # 1. GENERATE TXT FILE
    # -------------------------------------------------------------------------
    txt_content = """================================================================================
PROTOTYPE USER GUIDE & TESTING MANUAL
AI-Based Knowledge Retrieval Platform with Query Resolution System (Milestones 1-3)
================================================================================

--------------------------------------------------------------------------------
🚀 QUICK START GUIDE: HOW TO RUN THE PROTOTYPE
--------------------------------------------------------------------------------
1. Open your terminal in the project directory:
   Path: s:\\Personal\\My Projects\\AI-Knowledge-Retrieval-Platform

2. Install dependencies (if not already installed):
   Command: pip install -r requirements.txt

3. Launch the FastAPI Application Gateway Server:
   Command: python main.py

4. Open your Web Browser (Chrome / Edge / Firefox):
   URL: http://localhost:8000

5. Run Automated Test Suites anytime in terminal:
   - Milestone 1 Retrieval Evaluation:  python evaluation/evaluate_retrieval.py
   - Milestone 3 Verification Suite:     python evaluation/test_milestone3.py


================================================================================
SECTION 1: ABOUT THE PROTOTYPE & ARCHITECTURE
================================================================================
The AI Knowledge Retrieval Platform is an enterprise-grade multi-agent RAG engine designed for accurate, context-aware information extraction, voice interaction, and multi-domain query resolution.

Core Capabilities:
- Multi-Format Document Ingestion: Ingests PDF, DOCX, TXT, and CSV files.
- Decoupled 5-Agent Architecture:
  * Query Understanding Agent: Classifies queries (Factual, Procedural, Comparative, Ambiguous, Out-of-Bounds).
  * Retrieval Agent: Ranks vector similarity search and filters out low-confidence noise (< 0.30 score).
  * Clarification Agent: Detects ambiguity/multi-part requests and generates clarification follow-up prompts.
  * Response Generation Agent: Synthesizes grounded answers strictly from context with exact citations.
  * Conversation Memory Agent: Resolves coreference pronouns ('it', 'its', 'that policy') across turns.
- Web Speech API: Voice Speech-to-Text (STT) mic input & Text-to-Speech (TTS) playback.
- Response Transparency Panel: Displays retrieved evidence chunks, similarity meters, and metadata.


================================================================================
SECTION 2: FEATURE-BY-FEATURE USAGE GUIDE
================================================================================

FEATURE 1: RAG & WEB SPEECH QUERY INTERFACE
- Text Query Input: Type your question in the textarea and click "Submit Query".
- Voice STT (Speech-to-Text): Click the microphone icon -> Speak your question -> Transcribed text automatically appears in the box.
- Voice TTS (Text-to-Speech): Click "Play Voice" to listen to spoken answers with Play, Pause, Resume, and Stop controls.
- Domain Filter: Restrict search across Healthcare, Finance, or All domains.

FEATURE 2: CLASSIFICATION & CONFIDENCE BADGES
- Query Type Badge: Displays FACTUAL, PROCEDURAL, COMPARATIVE, or AMBIGUOUS.
- Routing Path Badge: Displays Retrieval Flow vs Clarification Flow.
- Application Confidence Badge: Displays HIGH CONFIDENCE, MEDIUM CONFIDENCE, or LOW CONFIDENCE with percentage score.

FEATURE 3: CLARIFICATION FEEDBACK & REFINEMENT LOOP
- Triggers when an ambiguous question like "how much?" is asked.
- Displays suggested question pills + refinement text input box.
- Type refinement or click a suggestion pill -> Submits to /api/query/refine and displays grounded answer.

FEATURE 4: MULTI-TURN CONVERSATION MEMORY & COREFERENCE RESOLUTION
- Turn 1: Ask "Tell me about Plan A medical coverage".
- Turn 2: Ask "What is its deductible?".
- Memory Agent automatically resolves "its" to "Plan A medical coverage" and returns the $500 deductible.

FEATURE 5: DOCUMENT INGESTION MODULE
- Drag & drop or browse .pdf, .docx, .txt, or .csv files up to 25MB.
- Set chunk size (default 500) and overlap (default 100).
- Click "Process & Index Document" -> View live terminal log output.
- Manage uploaded documents in the inventory table (view total chunks, delete document).

FEATURE 6: RESPONSE TRANSPARENCY PANEL
- Click to expand "Inspect Evidence & Source Transparency Panel".
- Views: Document Name, Section Heading, Page/Row Number, Similarity Percentage Meter, Chunk ID, and full text preview.

FEATURE 7: RETRIEVAL VALIDATION BENCHMARK DASHBOARD
- Click "Run Retrieval Evaluation Suite" button.
- View live accuracy meters for Top-1, Top-3, Top-5, and Rejection accuracy across Healthcare & Finance domains.


================================================================================
SECTION 3: STEP-BY-STEP TEST SCENARIOS & EXAMPLE QUERIES
================================================================================

TEST SCENARIO 1: FACTUAL QUERY LOOKUP
- Action: Type or speak "What is the waiting period for medical coverage eligibility?"
- Expected Classification: TYPE: FACTUAL | ROUTE: Retrieval Flow
- Expected Result: "Based on Healthcare Policy 2026: Employees become eligible for full comprehensive medical coverage after completing 30 days of continuous employment [1]."
- Confidence: HIGH CONFIDENCE (84.6%)
- Citation: [1] healthcare_policy.txt (Page/Row 1)

TEST SCENARIO 2: PROCEDURAL STEP-BY-STEP FORMATTING
- Action: Type or speak "What are the steps to submit an out of network claim?"
- Expected Classification: TYPE: PROCEDURAL | ROUTE: Retrieval Flow
- Expected Result: Formatted step-by-step instructions:
  1. Obtain itemized medical bill containing standard CPT/ICD-10 codes.
  2. Complete Form HC-104 within 90 calendar days of service.
  3. Attach original receipts and submit claim packet.
- Confidence: HIGH CONFIDENCE / MEDIUM CONFIDENCE

TEST SCENARIO 3: COMPARATIVE ANALYSIS
- Action: Type or speak "Compare coverage and deductible for Plan A versus Plan B"
- Expected Classification: TYPE: COMPARATIVE | ROUTE: Retrieval Flow
- Expected Result: Side-by-side comparison summary detailing Plan A ($500 deductible, 80/20 split) vs Plan B ($2,000 deductible + $1,000 HSA contribution).

TEST SCENARIO 4: MULTI-TURN COREFERENCE MEMORY
- Turn 1 Input: "Tell me about Plan A medical coverage"
- Turn 2 Input: "What is its deductible?"
- Expected Result: Memory Agent resolves "its" to "Plan A medical coverage" and returns the $500 deductible.

TEST SCENARIO 5: CLARIFICATION REFINEMENT LOOP
- Step 1 Input: "how much?"
- Step 1 Output: Clarification Bar appears with suggested question pills.
- Step 2 Input: Type "lodging reimbursement limit" in refinement box and click Submit Refinement.
- Expected Result: Refines query and returns $250 domestic lodging limit per night.

TEST SCENARIO 6: OUT-OF-BOUNDS ANTI-HALLUCINATION REJECTION
- Action: Type or speak "What is the company policy for pet insurance?"
- Expected Classification: TYPE: OUT_OF_BOUNDS | ROUTE: Clarification Flow
- Expected Result: "⚠️ The query falls outside the knowledge base scope. Please ask a question related to uploaded document policies."

TEST SCENARIO 7: DOCUMENT INGESTION WORKFLOW
- Action: Drag & drop a new .csv or .pdf file in Document Ingestion tab -> Click "Process & Index Document".
- Expected Result: Terminal log displays parsing, chunking, and indexing. Document appears in indexed table with chunk count.

TEST SCENARIO 8: AUTOMATED TEST SUITE VERIFICATION
- Action: Run "python evaluation/test_milestone3.py" in terminal.
- Expected Result: "=== ALL MILESTONE 3 TESTS PASSED SUCCESSFULLY! ==="
================================================================================
"""

    txt_path = "s:/Personal/My Projects/AI-Knowledge-Retrieval-Platform/Prototype_User_Guide_and_Testing_Manual.txt"
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(txt_content)
    print(f"Successfully generated {txt_path}")

    # -------------------------------------------------------------------------
    # 2. GENERATE DOCX FILE
    # -------------------------------------------------------------------------
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("PROTOTYPE USER GUIDE & TESTING MANUAL")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 58, 138)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("AI-Based Knowledge Retrieval Platform | Full System Prototype (Milestones 1-3)")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(37, 99, 235)
    
    doc.add_paragraph()

    # QUICK START BOX
    table_qs = doc.add_table(rows=1, cols=1)
    table_qs.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_qs = table_qs.cell(0, 0)
    set_cell_background(cell_qs, "EFF6FF")
    
    p_qs = cell_qs.paragraphs[0]
    r_qs_bold = p_qs.add_run("🚀 QUICK START: HOW TO RUN THE PROTOTYPE:\n")
    r_qs_bold.font.bold = True
    r_qs_bold.font.color.rgb = RGBColor(30, 58, 138)
    r_qs_text = p_qs.add_run(
        "1. Open Terminal in project directory: s:\\Personal\\My Projects\\AI-Knowledge-Retrieval-Platform\n"
        "2. Run command: python main.py\n"
        "3. Open Browser at: http://localhost:8000\n"
        "4. Automated Test Suites: python evaluation/test_milestone3.py"
    )
    r_qs_text.font.size = Pt(10.5)

    doc.add_paragraph()

    # SECTION 1: ABOUT THE PROTOTYPE
    h1 = doc.add_heading(level=1)
    h1.add_run("1. About the Prototype & Architecture").font.color.rgb = RGBColor(30, 58, 138)
    
    doc.add_paragraph(
        "The AI-Based Knowledge Retrieval Platform is an enterprise-grade multi-agent RAG engine that resolves complex queries "
        "across PDF, DOCX, TXT, and CSV documents with strict anti-hallucination guardrails, Web Speech voice interaction, "
        "and expandable evidence transparency."
    )

    capabilities = [
        ("Multi-Format Ingestion: ", "Parses PDF, DOCX, TXT, and CSV files into 500-character overlapping chunks."),
        ("5-Agent RAG Pipeline: ", "Query Understanding, Retrieval, Response Generation, Clarification, and Conversation Memory Agents."),
        ("Web Speech API Integration: ", "Native Speech-to-Text (STT) mic recording and Text-to-Speech (TTS) voice playback controls."),
        ("Response Transparency Panel: ", "Expandable evidence inspector displaying relevance meters, metadata, and citation mappings.")
    ]
    for title, desc in capabilities:
        bp = doc.add_paragraph(style='List Bullet')
        r_bt = bp.add_run(title)
        r_bt.font.bold = True
        r_bt.font.color.rgb = RGBColor(37, 99, 235)
        bp.add_run(desc)

    doc.add_paragraph()

    # SECTION 2: TEST SCENARIOS TABLE
    h2 = doc.add_heading(level=1)
    h2.add_run("2. Step-by-Step Test Scenarios & Example Queries").font.color.rgb = RGBColor(30, 58, 138)

    test_scenarios = [
        {
            "num": "1",
            "name": "Factual Query Lookup",
            "input": '"What is the waiting period for medical coverage eligibility?"',
            "result": "Returns 30 days eligibility answer with page citation [1] healthcare_policy.txt. Confidence: HIGH (84.6%)."
        },
        {
            "num": "2",
            "name": "Procedural Step Formatting",
            "input": '"What are the steps to submit an out of network claim?"',
            "result": "Formats output into numbered step-by-step instructions (1. Itemized bill, 2. Form HC-104, 3. Receipts)."
        },
        {
            "num": "3",
            "name": "Comparative Analysis",
            "input": '"Compare coverage and deductible for Plan A versus Plan B"',
            "result": "Formats side-by-side comparison detailing Plan A ($500 deductible) vs Plan B ($2,000 deductible + $1,000 HSA)."
        },
        {
            "num": "4",
            "name": "Multi-Turn Coreference Memory",
            "input": 'Turn 1: "Tell me about Plan A medical coverage"\nTurn 2: "What is its deductible?"',
            "result": "Memory Agent resolves 'its' to 'Plan A medical coverage' and returns $500 deductible."
        },
        {
            "num": "5",
            "name": "Clarification Refinement Loop",
            "input": 'Step 1: "how much?"\nStep 2: Type "lodging reimbursement limit"',
            "result": "Clarification Bar appears -> Refines query -> Returns $250 domestic lodging limit per night."
        },
        {
            "num": "6",
            "name": "Out-of-Bounds Rejection",
            "input": '"What is the company policy for pet insurance?"',
            "result": "Safely rejects off-topic query with zero AI hallucination."
        }
    ]

    for ts in test_scenarios:
        table_t = doc.add_table(rows=3, cols=2)
        table_t.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        cell_h = table_t.cell(0, 0)
        cell_h.merge(table_t.cell(0, 1))
        set_cell_background(cell_h, "1E3A8A")
        p_h = cell_h.paragraphs[0]
        r_h = p_h.add_run(f"TEST SCENARIO {ts['num']}: {ts['name']}")
        r_h.font.bold = True
        r_h.font.color.rgb = RGBColor(255, 255, 255)
        
        table_t.cell(1, 0).paragraphs[0].add_run("Inputs / Action:").font.bold = True
        table_t.cell(1, 1).paragraphs[0].add_run(ts["input"]).font.color.rgb = RGBColor(37, 99, 235)
        
        table_t.cell(2, 0).paragraphs[0].add_run("Expected Output:").font.bold = True
        table_t.cell(2, 1).paragraphs[0].add_run(ts["result"])
        
        doc.add_paragraph()

    docx_path = "s:/Personal/My Projects/AI-Knowledge-Retrieval-Platform/Prototype_User_Guide_and_Testing_Manual.docx"
    doc.save(docx_path)
    print(f"Successfully generated {docx_path}")

if __name__ == "__main__":
    create_prototype_guide_files()
